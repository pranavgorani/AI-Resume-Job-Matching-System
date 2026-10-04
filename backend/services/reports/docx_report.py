import io
import datetime
from typing import Any, Dict
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def generate_docx_report(data: Dict[str, Any]) -> bytes:
    """
    Generates a professional DOCX Word document report using python-docx.
    """
    doc = Document()

    # Set Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # -------------------------------------------------------------
    # COVER / HEADER
    # -------------------------------------------------------------
    p_brand = doc.add_paragraph()
    r_brand = p_brand.add_run("TALENTPROOF AI")
    r_brand.font.size = Pt(12)
    r_brand.font.bold = True
    r_brand.font.color.rgb = RGBColor(79, 70, 229) # Indigo

    p_title = doc.add_paragraph()
    r_title = p_title.add_run("Recruitment Intelligence Analytics Report")
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("Evidence-Based Candidate Matching & Risk Intelligence")
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph(f"Date: {data.get('date_generated', datetime.datetime.utcnow().strftime('%Y-%m-%d'))}  |  Position: {data.get('job_title', 'All Positions')}  |  Candidates Analyzed: {len(data.get('candidates', []))}")
    doc.add_paragraph("─" * 60)

    # -------------------------------------------------------------
    # 1. EXECUTIVE SUMMARY
    # -------------------------------------------------------------
    h1 = doc.add_heading("1. Executive Summary", level=1)
    h1.style.font.color.rgb = RGBColor(30, 41, 59)
    kpis = data.get("kpis", {})

    p_exec = doc.add_paragraph(
        f"This executive intelligence report synthesizes evaluations for {kpis.get('total_candidates', 0)} candidates "
        f"applying for {data.get('job_title', 'the specified position')}. TalentProof AI eliminates resume keyword fraud "
        f"by constructing verifiable Evidence Graphs linking requirements directly to verified project and chronological accomplishments. "
        f"Across the active pool, candidates scored an average match rate of {kpis.get('average_match', 0)}% and evidence confidence of "
        f"{kpis.get('average_evidence', 0)}%. Currently, {kpis.get('shortlist_rate', 0)}% of candidates demonstrate strong hiring alignment."
    )
    p_exec.style.font.size = Pt(10.5)

    # -------------------------------------------------------------
    # 2. KPI TABLE
    # -------------------------------------------------------------
    h2 = doc.add_heading("2. Key Performance Indicators", level=1)
    h2.style.font.color.rgb = RGBColor(30, 41, 59)

    kpi_table = doc.add_table(rows=5, cols=2)
    kpi_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    kpi_rows = [
        ("Total Candidates Evaluated", f"{kpis.get('total_candidates', 0)}"),
        ("Average Match Score", f"{kpis.get('average_match', 0)}%"),
        ("Average Evidence Score", f"{kpis.get('average_evidence', 0)}%"),
        ("Hiring Confidence", f"{kpis.get('average_hiring_confidence', 0)}%"),
        ("Verification Required Rate", f"{kpis.get('verification_rate', 0)}% ({kpis.get('high_risk_candidates', 0)} high risk)")
    ]

    for idx, (label, val) in enumerate(kpi_rows):
        kpi_table.cell(idx, 0).text = label
        kpi_table.cell(idx, 1).text = val
        kpi_table.cell(idx, 0).paragraphs[0].runs[0].font.bold = True

    # -------------------------------------------------------------
    # 3. CANDIDATE RANKING TABLE
    # -------------------------------------------------------------
    h3 = doc.add_heading("3. Candidate Ranking Table", level=1)
    h3.style.font.color.rgb = RGBColor(30, 41, 59)

    candidates = data.get("candidates", [])[:15]
    table = doc.add_table(rows=len(candidates) + 1, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["Rank", "Candidate Name", "Match", "Evidence", "Experience", "Risk Level", "Recommendation"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True

    for row_idx, c in enumerate(candidates, 1):
        table.cell(row_idx, 0).text = f"#{row_idx}"
        table.cell(row_idx, 1).text = c.get("candidate_name", "")
        table.cell(row_idx, 2).text = f"{c.get('match_score', 0)}%"
        table.cell(row_idx, 3).text = f"{c.get('evidence_score', 0)}%"
        table.cell(row_idx, 4).text = f"{c.get('experience', 0)} yrs"
        table.cell(row_idx, 5).text = c.get("risk_level", "LOW")
        table.cell(row_idx, 6).text = c.get("recommendation", "CONSIDER").replace("_", " ")

    # -------------------------------------------------------------
    # 4. MATCH & EVIDENCE SCORE ANALYSIS
    # -------------------------------------------------------------
    h4 = doc.add_heading("4. Match & Evidence Score Analysis", level=1)
    h4.style.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "Match scores are calculated using the 7-factor deterministic engine: "
        "35% Required Skills, 20% Relevant Experience, 15% Project Evidence, "
        "10% Education/Certifications, 10% Preferred Skills, 5% Domain Knowledge, and 5% Evidence Confidence. "
        "Scores are not subject to generative LLM hallucinations."
    )

    # -------------------------------------------------------------
    # 5. SKILL ANALYSIS
    # -------------------------------------------------------------
    h5 = doc.add_heading("5. Skill & Transferability Analysis", level=1)
    h5.style.font.color.rgb = RGBColor(30, 41, 59)
    top_skills = data.get("distributions", {}).get("top_skills", [])
    if top_skills:
        for s in top_skills[:6]:
            doc.add_paragraph(f"• {s.get('skill')}: Verified in {s.get('count')} candidate portfolio(s)")
    else:
        doc.add_paragraph("No candidate skills recorded for this view.")

    # -------------------------------------------------------------
    # 6. RISK & CONTRADICTION ANALYSIS
    # -------------------------------------------------------------
    h6 = doc.add_heading("6. Risk & Contradiction Analysis", level=1)
    h6.style.font.color.rgb = RGBColor(30, 41, 59)
    risk_d = data.get("distributions", {}).get("risk_distribution", {})
    doc.add_paragraph(
        f"Timeline overlaps and unbacked skill assertions are catalogued with neutral recommendations. "
        f"Current breakdown: No Risk: {risk_d.get('No Risk', 0)}, Low Risk: {risk_d.get('Low', 0)}, "
        f"Medium Risk: {risk_d.get('Medium', 0)}, High Risk: {risk_d.get('High', 0)}."
    )

    # -------------------------------------------------------------
    # 7. HIRING FUNNEL
    # -------------------------------------------------------------
    h7 = doc.add_heading("7. Hiring Pipeline Funnel", level=1)
    h7.style.font.color.rgb = RGBColor(30, 41, 59)
    funnel = data.get("distributions", {}).get("hiring_funnel", [])
    for f in funnel:
        doc.add_paragraph(f"• {f.get('stage')}: {f.get('count')} candidates")

    # -------------------------------------------------------------
    # 8. AI RECRUITMENT INSIGHTS
    # -------------------------------------------------------------
    h8 = doc.add_heading("8. AI Recruitment Insights", level=1)
    h8.style.font.color.rgb = RGBColor(30, 41, 59)
    insights = data.get("ai_insights", [])
    for ins in insights:
        doc.add_paragraph(f"• {ins}")

    # -------------------------------------------------------------
    # 9. VERIFICATION REQUIRED CANDIDATES
    # -------------------------------------------------------------
    h9 = doc.add_heading("9. Verification Required Candidates", level=1)
    h9.style.font.color.rgb = RGBColor(30, 41, 59)
    verif = [c for c in data.get("candidates", []) if c.get("verification_status") == "REQUIRED"]
    if verif:
        for vc in verif:
            flags_str = "; ".join(vc.get("risk_flags_details", [])) or "Timeline check needed"
            doc.add_paragraph(f"• {vc.get('candidate_name')}: {flags_str}")
    else:
        doc.add_paragraph("✓ All candidate profiles verified with zero pending discrepancies.")

    # -------------------------------------------------------------
    # 10. RECOMMENDATIONS & CONCLUSION
    # -------------------------------------------------------------
    h10 = doc.add_heading("10. Recommendations & Conclusion", level=1)
    h10.style.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "1. Proceed to interview scheduling with Strongly Recommended candidates.\n"
        "2. Utilize the AI Interview Intelligence module for targeted gap verification.\n"
        "3. Fair Match Compliance is actively preserved; scoring models remain deterministic."
    )

    doc_stream = io.BytesIO()
    doc.save(doc_stream)
    return doc_stream.getvalue()
