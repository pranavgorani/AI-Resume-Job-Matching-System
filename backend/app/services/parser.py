import os
import logging
from pathlib import Path
from typing import Tuple

logger = logging.getLogger(__name__)

import io

def extract_text_from_bytes(file_bytes: bytes, filename: str) -> Tuple[str, str]:
    """
    In-memory text extractor from bytes. Does not require local filesystem persistence.
    Ideal for serverless runtime environments like Vercel.
    """
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return _extract_from_pdf_bytes(file_bytes), "pdf"
    elif ext in [".docx", ".doc"]:
        return _extract_from_docx_bytes(file_bytes), "docx"
    elif ext in [".txt", ".md", ".json"]:
        return _extract_from_txt_bytes(file_bytes), "txt"
    else:
        try:
            return _extract_from_txt_bytes(file_bytes), "txt"
        except Exception:
            raise ValueError(f"Unsupported file format: {ext}")

def extract_text_from_file(file_path: Path) -> Tuple[str, str]:
    """
    Extracts text from PDF, DOCX, or TXT.
    Returns (raw_text, detected_type).
    """
    try:
        with open(file_path, "rb") as f:
            data = f.read()
        return extract_text_from_bytes(data, file_path.name)
    except Exception as e:
        if isinstance(e, ValueError):
            raise
        raise ValueError(f"Could not read document: {str(e)}")

def _extract_from_pdf_bytes(file_bytes: bytes) -> str:
    try:
        from pypdf import PdfReader
        stream = io.BytesIO(file_bytes)
        reader = PdfReader(stream)
        text_parts = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        extracted = "\n".join(text_parts).strip()
        if not extracted:
            raise ValueError("PDF document appears empty or scanned without an extractable text layer.")
        return extracted
    except Exception as e:
        logger.error(f"Error parsing PDF: {e}")
        if isinstance(e, ValueError):
            raise
        raise ValueError(f"Failed to parse PDF document: {str(e)}")

def _extract_from_docx_bytes(file_bytes: bytes) -> str:
    try:
        import docx
        stream = io.BytesIO(file_bytes)
        doc = docx.Document(stream)
        full_text = []
        for para in doc.paragraphs:
            if para.text:
                full_text.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        full_text.append(cell.text)
        extracted = "\n".join(full_text).strip()
        if not extracted:
            raise ValueError("DOCX document appears empty.")
        return extracted
    except Exception as e:
        logger.error(f"Error parsing DOCX: {e}")
        if isinstance(e, ValueError):
            raise
        raise ValueError(f"DOCX parsing failed: {str(e)}. Please retry or upload as PDF.")

def _extract_from_txt_bytes(file_bytes: bytes) -> str:
    try:
        return file_bytes.decode("utf-8", errors="replace").strip()
    except Exception as e:
        logger.error(f"Error reading text bytes: {e}")
        raise ValueError(f"Failed to read text file: {str(e)}")

def _extract_from_pdf(file_path: Path) -> str:
    with open(file_path, "rb") as f:
        return _extract_from_pdf_bytes(f.read())

def _extract_from_docx(file_path: Path) -> str:
    with open(file_path, "rb") as f:
        return _extract_from_docx_bytes(f.read())

def _extract_from_txt(file_path: Path) -> str:
    with open(file_path, "rb") as f:
        return _extract_from_txt_bytes(f.read())

